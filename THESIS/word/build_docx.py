"""
Builds the Word thesis from the MTC thesis template and the converted
LaTeX content (prep_latex.py + pandoc).

  python THESIS/word/build_docx.py <template.docx> <conv_dir> <out.docx> [indexes.json]

conv_dir holds prep/{body,abstract,acknowledgments,snippets}.docx,
prep/side.json, refs.docx (IEEE bibliography in citation order) and
fig/*.png.
"""
from __future__ import annotations

import copy
import json
import re
import sys
import zipfile
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Cm
from lxml import etree

THESIS = Path(__file__).resolve().parents[1]
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
TEXT_W = 9971          # twips between the margins of the template (21.59 cm - 2.5 cm - 1.5 cm)
TABLE_W = 9900

TITLE = "Training-Free Adaptive Control of Detector and Tracker Operating Points for Multi-Object Tracking"
AUTHOR = "Ahmed Gouda Ismail"
SUP1, SUP1_AFF = "Mohamed S. Mohamed", "Military Technical College"
SUP2, SUP2_AFF = "Tarek Ahmed Mahmoud", "Egypt University of Informatics"
FIELD = "Computer Engineering"
YEAR = "2026"
AR_TITLE = "التحكم التكيفي دون تدريب في نقاط تشغيل الكاشف والمتتبع لتتبع الأجسام المتعددة"
AR_AUTHOR = "أحمد جودة إسماعيل"
AR_SUP1, AR_SUP2 = "محمد س. محمد", "طارق أحمد محمود"
AR_SUP2_AFF = "جامعة مصر للمعلوماتية"
AR_DEGREE = "للحصول على درجة الماجستير فى هندسة الحاسبات"


def el(tag, **attrs):
    e = etree.Element(qn(tag))
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def sub(parent, tag, **attrs):
    e = el(tag, **attrs)
    parent.append(e)
    return e


def text_of(e):
    return "".join(t.text or "" for t in e.iter(qn("w:t")))


# ---------------------------------------------------------------- pandoc input
def pandoc_paragraphs(path):
    x = zipfile.ZipFile(path).read("word/document.xml")
    root = etree.fromstring(x)
    body = root.find(qn("w:body"))
    return [p for p in body if p.tag == qn("w:p")]


def clean_runs(p):
    """Run-level content of a pandoc paragraph, ready for the template."""
    out = []

    def add_run(r):
        r = copy.deepcopy(r)
        rpr = r.find(qn("w:rPr"))
        if rpr is not None:
            st = rpr.find(qn("w:rStyle"))
            if st is not None:
                if st.get(qn("w:val")) == "VerbatimChar":
                    f = el("w:rFonts", **{"w:ascii": "Courier New", "w:hAnsi": "Courier New", "w:cs": "Courier New"})
                    rpr.insert(0, f)
                rpr.remove(st)
        for t in list(r.findall(qn("w:t"))):
            if t.text and "\t" in t.text:
                parts = t.text.split("\t")
                idx = list(r).index(t)
                r.remove(t)
                for k, part in enumerate(parts):
                    if k:
                        r.insert(idx, el("w:tab"))
                        idx += 1
                    if part:
                        nt = el("w:t")
                        nt.text = part
                        nt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                        r.insert(idx, nt)
                        idx += 1
        out.append(r)
    for c in p:
        if c.tag == qn("w:r"):
            add_run(c)
        elif c.tag == qn("w:hyperlink"):
            for r in c.findall(qn("w:r")):
                add_run(r)
        elif c.tag == f"{{{M_NS}}}oMath":
            out.append(copy.deepcopy(c))
    return out


def strip_prefix(runs, n):
    """Remove the first n characters of text from a run list."""
    for r in runs:
        if n <= 0:
            break
        if r.tag != qn("w:r"):
            continue
        for t in r.findall(qn("w:t")):
            if n <= 0:
                break
            s = t.text or ""
            k = min(n, len(s))
            t.text = s[k:]
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            n -= k
    return [r for r in runs if not (r.tag == qn("w:r") and r.find(qn("w:t")) is not None
                                     and all(not (t.text or "") for t in r.findall(qn("w:t")))
                                     and r.find(qn("w:tab")) is None and r.find(qn("w:br")) is None)]


RPR_ORDER = {k: i for i, k in enumerate(
    "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish "
    "webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang "
    "eastAsianLayout specVanish oMath rPrChange".split())}
PPR_ORDER = {k: i for i, k in enumerate(
    "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
    "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
    "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
    "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange".split())}

MARK = re.compile(r"^QQ(CH|HA|HB|HC|FIG|TAB|EQ|ALG|LI|DT|DD|SN)(.*?)QQ ?")


def marker(p):
    m = MARK.match(text_of(p))
    return (m.group(1), m.group(2), m.end()) if m else (None, None, 0)


def run(text, bold=False, italic=False, size=None, rtl=False, base_rpr=None):
    r = el("w:r")
    rpr = copy.deepcopy(base_rpr) if base_rpr is not None else el("w:rPr")
    if bold:
        sub(rpr, "w:b")
        sub(rpr, "w:bCs")
    if italic:
        sub(rpr, "w:i")
    if size:
        sub(rpr, "w:sz", **{"w:val": size})
        sub(rpr, "w:szCs", **{"w:val": size})
    if rtl and rpr.find(qn("w:rtl")) is None:
        sub(rpr, "w:rtl")
    if len(rpr):
        r.append(rpr)
    t = sub(r, "w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return r


def tab_run():
    r = el("w:r")
    sub(r, "w:tab")
    return r


def set_size(runs, size):
    for r in runs:
        if r.tag == f"{{{M_NS}}}oMath":
            for mr in r.iter(f"{{{M_NS}}}r"):
                wr = mr.find(qn("w:rPr"))
                if wr is None:
                    wr = el("w:rPr")
                    mrp = mr.find(f"{{{M_NS}}}rPr")
                    mr.insert(1 if mrp is not None else 0, wr)
                for tag in ("w:sz", "w:szCs"):
                    old = wr.find(qn(tag))
                    if old is not None:
                        wr.remove(old)
                    sub(wr, tag, **{"w:val": size})
            continue
        if r.tag != qn("w:r"):
            continue
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = el("w:rPr")
            r.insert(0, rpr)
        for tag in ("w:sz", "w:szCs"):
            old = rpr.find(qn(tag))
            if old is not None:
                rpr.remove(old)
            sub(rpr, tag, **{"w:val": size})
    return runs


def bold_runs(runs):
    for r in runs:
        if r.tag != qn("w:r"):
            continue
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = el("w:rPr")
            r.insert(0, rpr)
        if rpr.find(qn("w:b")) is None:
            rpr.insert(0, el("w:b"))
    return runs


def split_sentence(runs):
    """Split a run list after its first sentence (caption title / description)."""
    full = "".join((t.text or "") for r in runs if r.tag == qn("w:r") for t in r.findall(qn("w:t")))
    depth, cut = 0, None
    for k, ch in enumerate(full):
        if ch in "([":
            depth += 1
        elif ch in ")]":
            depth -= 1
        elif ch == "." and depth == 0 and k + 1 < len(full) and full[k + 1] == " ":
            word = re.findall(r"[A-Za-z]+$", full[:k])
            if word and word[-1] in ("e", "g", "i", "vs", "Fig", "Eq", "Sec", "al", "approx", "cf", "Tab"):
                continue
            if k > 0 and full[k - 1].isdigit() and k + 2 < len(full) and full[k + 2].isdigit():
                continue
            cut = k + 1
            break
    if cut is None:
        return runs, []
    first, rest, pos = [], [], 0
    for r in runs:
        if r.tag != qn("w:r") or not r.findall(qn("w:t")):
            (first if pos < cut else rest).append(r)
            continue
        txt = "".join(t.text or "" for t in r.findall(qn("w:t")))
        if pos + len(txt) <= cut:
            first.append(r)
        elif pos >= cut:
            rest.append(r)
        else:
            a, b = copy.deepcopy(r), copy.deepcopy(r)
            k = cut - pos
            for t in a.findall(qn("w:t")):
                t.text = (t.text or "")[:k]
                k = max(0, k - len(t.text))
            k = cut - pos
            for t in b.findall(qn("w:t")):
                s = t.text or ""
                t.text = s[k:] if k < len(s) else ""
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                k = max(0, k - len(s))
            first.append(a)
            rest.append(b)
        pos += len(txt)
    rest = strip_prefix(rest, 1) if rest else rest
    return first, rest


# ---------------------------------------------------------------- template
class Builder:
    def __init__(self, template, conv, indexes=None):
        self.indexes = indexes
        self.d = docx.Document(template)
        self.conv = Path(conv)
        self.body = self.d.element.body
        self.side = json.loads((self.conv / "prep/side.json").read_text())
        snip = pandoc_paragraphs(self.conv / "prep/snippets.docx")
        self.snip = {}
        for p in snip:
            kind, sid, n = marker(p)
            if kind == "SN":
                self.snip[sid] = strip_prefix(clean_runs(p), n)
        kids = list(self.body)
        self.proto = {k: copy.deepcopy(kids[i]) for k, i in dict(body=170, h1=168, h2=626, fig=174, tab=264,
                                                                     ref=662, abbr=143, absb=47).items()}

    def snippet(self, sid):
        return [copy.deepcopy(r) for r in self.snip.get(sid, [])] if sid else []

    # paragraphs -------------------------------------------------------------
    def para(self, runs, style="NoSpacing", first_line=None, left=None, hanging=None, jc=None, after=240,
             before=None, keep_next=False, tabs=None, page_break_before=False, line=None):
        p = el("w:p")
        ppr = sub(p, "w:pPr")
        sub(ppr, "w:pStyle", **{"w:val": style})
        if keep_next:
            sub(ppr, "w:keepNext")
        if page_break_before:
            sub(ppr, "w:pageBreakBefore")
        if tabs:
            t = sub(ppr, "w:tabs")
            for kind, pos in tabs:
                sub(t, "w:tab", **{"w:val": kind, "w:pos": pos})
        sp = {"w:after": after}
        if before is not None:
            sp["w:before"] = before
        if line is not None:
            sp.update({"w:line": line, "w:lineRule": "auto"})
        sub(ppr, "w:spacing", **sp)
        ind = {}
        if left is not None:
            ind["w:left"] = left
        if hanging is not None:
            ind["w:hanging"] = hanging
        if first_line is not None:
            ind["w:firstLine"] = first_line
        if ind:
            sub(ppr, "w:ind", **ind)
        if jc:
            sub(ppr, "w:jc", **{"w:val": jc})
        for r in runs:
            if style in ("Normal", "ListParagraph") and r.tag == qn("w:r"):
                rpr = r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = el("w:rPr")
                    r.insert(0, rpr)
                if rpr.find(qn("w:rFonts")) is None:
                    rpr.insert(0, el("w:rFonts", **{"w:asciiTheme": "majorBidi", "w:hAnsiTheme": "majorBidi",
                                                    "w:cstheme": "majorBidi"}))
            p.append(r)
        return p

    def page_break(self):
        p = el("w:p")
        r = sub(p, "w:r")
        sub(r, "w:br", **{"w:type": "page"})
        return p

    def chapter_heading(self, label, title_runs):
        p = el("w:p")
        ppr = sub(p, "w:pPr")
        sub(ppr, "w:pStyle", **{"w:val": "Heading1"})
        sub(ppr, "w:pageBreakBefore")
        p.append(run(label))
        r = el("w:r")
        sub(r, "w:br")
        p.append(r)
        for x in title_runs:
            p.append(x)
        return p

    def heading(self, level, num, title_runs):
        p = el("w:p")
        ppr = sub(p, "w:pPr")
        sub(ppr, "w:pStyle", **{"w:val": "Heading2" if level == 2 else "Heading3"})
        sub(ppr, "w:keepNext")
        npr = sub(ppr, "w:numPr")
        sub(npr, "w:ilvl", **{"w:val": 0})
        sub(npr, "w:numId", **{"w:val": 0})
        w = 576 if level == 2 else 720
        sub(ppr, "w:ind", **{"w:left": w, "w:hanging": w})
        if num:
            p.append(run(num + " "))
        for x in title_runs:
            p.append(x)
        return p

    def picture(self, path, width_cm):
        tmp = self.d.add_paragraph()
        tmp.add_run().add_picture(str(path), width=Cm(width_cm))
        self.body.remove(tmp._p)
        r = tmp._p.find(qn("w:r"))
        return self.para([r], first_line=0, jc="center", keep_next=True, before=120, after=60, line=240)

    def caption(self, kind, num, runs):
        """kind 'fig' / 'tab' -> (caption paragraph, description paragraph or None)."""
        title, rest = split_sentence(runs)
        label = f"Fig. ({num}) " if kind == "fig" else f"Table ({num}) "
        style = "myfiguresibrahim" if kind == "fig" else "mytablesibrahim"
        cap = self.para([run(label)] + title, style=style, after=60 if rest else 200, keep_next=kind == "tab" or bool(rest),
                        before=0 if kind == "fig" else 200)
        desc = None
        if rest:
            desc = self.para(set_size(rest, 20), first_line=0, jc="both", after=200 if kind == "fig" else 80,
                             keep_next=kind == "tab", line=240)
        return cap, desc

    def equation(self, num, omath):
        runs = [tab_run(), omath, tab_run()]
        if num:
            runs.append(run(f"({num})"))
        return self.para(runs, first_line=0, tabs=[("center", TEXT_W // 2), ("right", TEXT_W)], after=120, before=120)

    # tables -----------------------------------------------------------------
    def table(self, t):
        cols = t["cols"]
        ncol = len(cols)
        rows = t["rows"]
        size = 16 if (ncol >= 7 or t.get("sideways")) else 18
        widths = []
        for k, c in enumerate(cols):
            if c["width"]:
                m = re.match(r"([0-9.]+)\s*(cm|mm|in|pt|em)?", c["width"])
                val = float(m.group(1)) * {"cm": 1, "mm": 0.1, "in": 2.54, "pt": 0.0353, "em": 0.42, None: 1}[m.group(2)]
                widths.append(val)
            else:
                longest = 4
                for r in rows:
                    cells = r["cells"]
                    pos = 0
                    for c2 in cells:
                        if pos == k and c2["span"] == 1 and c2["sid"]:
                            txt = "".join(text_of(x) for x in self.snip.get(c2["sid"], []) if x.tag == qn("w:r"))
                            longest = max(longest, min(len(txt), 40))
                        pos += c2["span"]
                widths.append(longest * 0.19 + 0.3)
        scale = TABLE_W / (sum(widths) * 567)
        tw = [max(400, int(w * 567 * scale)) for w in widths]
        tw[-1] += TABLE_W - sum(tw)
        tbl = el("w:tbl")
        tpr = sub(tbl, "w:tblPr")
        sub(tpr, "w:tblStyle", **{"w:val": "TableGrid"})
        sub(tpr, "w:tblW", **{"w:w": TABLE_W, "w:type": "dxa"})
        sub(tpr, "w:jc", **{"w:val": "center"})
        sub(tpr, "w:tblLayout", **{"w:type": "fixed"})
        mar = sub(tpr, "w:tblCellMar")
        sub(mar, "w:left", **{"w:w": 57, "w:type": "dxa"})
        sub(mar, "w:right", **{"w:w": 57, "w:type": "dxa"})
        sub(tpr, "w:tblLook", **{"w:val": "04A0", "w:firstRow": 1, "w:lastRow": 0, "w:firstColumn": 0,
                                 "w:lastColumn": 0, "w:noHBand": 1, "w:noVBand": 1})
        grid = sub(tbl, "w:tblGrid")
        for w in tw:
            sub(grid, "w:gridCol", **{"w:w": w})
        header_rows = 1 if len(rows) > 1 and (len(rows) < 2 or rows[1]["rule_before"]) else 0
        for ri, r in enumerate(rows):
            tr = sub(tbl, "w:tr")
            trpr = sub(tr, "w:trPr")
            sub(trpr, "w:cantSplit")
            if ri < header_rows:
                sub(trpr, "w:tblHeader")
            pos = 0
            for c in r["cells"]:
                tc = sub(tr, "w:tc")
                tcpr = sub(tc, "w:tcPr")
                span = min(c["span"], ncol - pos)
                sub(tcpr, "w:tcW", **{"w:w": sum(tw[pos:pos + span]), "w:type": "dxa"})
                if span > 1:
                    sub(tcpr, "w:gridSpan", **{"w:val": span})
                if ri < header_rows:
                    sub(tcpr, "w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": "D9D9D9"})
                sub(tcpr, "w:vAlign", **{"w:val": "center"})
                align = c["align"] or (cols[pos]["align"] if pos < ncol else "l")
                if span == ncol and c["align"] is None:
                    align = "l"
                runs = set_size(self.snippet(c["sid"]), size)
                if ri < header_rows:
                    bold_runs(runs)
                p = self.para(runs, style="Normal", first_line=0, after=0, line=240,
                              jc={"l": "left", "c": "center", "r": "right"}[align])
                tc.append(p)
                pos += span
            while pos < ncol:
                tc = sub(tr, "w:tc")
                tcpr = sub(tc, "w:tcPr")
                sub(tcpr, "w:tcW", **{"w:w": tw[pos], "w:type": "dxa"})
                tc.append(self.para([], style="Normal", after=0, line=240))
                pos += 1
        return tbl

    # content ----------------------------------------------------------------
    def content(self):
        out = []
        paras = pandoc_paragraphs(self.conv / "prep/body.docx")
        i = 0
        after_eq = False
        while i < len(paras):
            p = paras[i]
            kind, key, n = marker(p)
            runs = strip_prefix(clean_runs(p), n) if kind else clean_runs(p)
            if kind == "CH":
                ch = self.side["chapters"][key]
                label = f"Appendix {key}" if key.isalpha() else f"Chapter {key}"
                title = self.snip_from_latex_title(ch["title"])
                out.append(self.chapter_heading(label, title))
            elif kind in ("HA", "HB", "HC"):
                out.append(self.heading(2 if kind == "HA" else 3, key, runs))
            elif kind == "FIG":
                f = self.side["figures"][key]
                img = self.conv / "fig" / ("fig" + key.replace(".", "_") + ".png")
                width = min(16.0, 16.0 * f["src"]["width"]) if "pdf" in f["src"] else 16.0
                out.append(self.picture(img, width))
                cap, desc = self.caption("fig", key, self.snippet(f["caption"]))
                out.append(cap)
                if desc is not None:
                    out.append(desc)
            elif kind == "ALG":
                a = self.side["algos"][key]
                out.append(self.picture(self.conv / "fig" / f"alg{key}.png", 15.0))
                title, rest = split_sentence(self.snippet(a["caption"]))
                out.append(self.para(bold_runs([run(f"Algorithm {key}. ")] + title + rest), style="Normal",
                                     jc="center", after=240, first_line=0, line=240))
            elif kind == "TAB":
                t = self.side["tables"][key]
                if t["caption"]:
                    cap, desc = self.caption("tab", key, self.snippet(t["caption"]))
                    out.append(cap)
                    if desc is not None:
                        out.append(desc)
                out.append(self.table(t))
                out.append(self.para([], after=120))
            elif kind == "EQ":
                i += 1
                mp = paras[i]
                omp = mp.find(f"{{{M_NS}}}oMathPara")
                om = copy.deepcopy(omp.find(f"{{{M_NS}}}oMath"))
                out.append(self.equation(key, om))
                after_eq = True
                i += 1
                continue
            elif kind == "LI":
                mark = self.side["lists"][key]["mark"]
                out.append(self.para([run(mark), tab_run()] + runs, first_line=None, left=567, hanging=567, after=120))
            elif kind == "DT":
                out.append(self.para(runs, first_line=0, after=0, keep_next=True))
            elif kind == "DD":
                out.append(self.para(runs, first_line=0, left=0, after=200))
            else:
                if runs:
                    out.append(self.para(runs, first_line=0 if after_eq else None))
            after_eq = False
            i += 1
        return out

    def snip_from_latex_title(self, title):
        return [run(re.sub(r"\\[a-z]+\{([^}]*)\}", r"\1", title))]

    # front matter -----------------------------------------------------------
    def replace_paragraph_text(self, p, new):
        ts = list(p.iter(qn("w:t")))
        ts[0].text = new
        ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        for t in ts[1:]:
            t.text = ""

    def front(self):
        kids = list(self.body)
        texts = {i: text_of(k).strip() for i, k in enumerate(kids)}
        rep = {
            "Malicious Behavior Detection for Safety Applications in Vehicular Network": TITLE,
            "Capt. Eng. Ibrahim Bassiony Elsaka": AUTHOR,
            "Maj. Gen. (R.) Prof. Dr. Gouda Ismail Salama": SUP1,
            "Col. Dr. Sherif Ibrahim Morsy": SUP2,
            "Cairo 2024": f"Cairo {YEAR}",
            "Ibrahim Bassiony": AUTHOR,
            "الكشف عن هجمات تزوير الموقع في شبكات المركبات": AR_TITLE,
            "نقيب مهندس / ابراهيم بسيونى السقا": AR_AUTHOR,
            "للحصول على درجة الماجستير فى الهندسة الكهربية": AR_DEGREE,
            "لواء أ.د./ جوده اسماعيل سلامه": AR_SUP1,
            "عقيد د./ شريف ابراهيم مرسى": AR_SUP2,
            "القاهرة 4202": f"القاهرة {YEAR}",
        }
        for i, k in enumerate(kids):
            if k.tag != qn("w:p"):
                continue
            t = texts[i]
            if t in rep:
                self.replace_paragraph_text(k, rep[t])
                if t == "Col. Dr. Sherif Ibrahim Morsy":
                    self.replace_paragraph_text(kids[i + 1], SUP2_AFF)
                if t == "عقيد د./ شريف ابراهيم مرسى":
                    self.replace_paragraph_text(kids[i + 1], AR_SUP2_AFF)
            for tt in k.iter(qn("w:t")):
                if tt.text == "Electrical Engineering":
                    tt.text = FIELD
            if t.startswith("Date of Discussion"):
                for tt in k.iter(qn("w:t")):
                    if tt.text == "2024":
                        tt.text = YEAR
        # committee table: names and ranks are filled in by the department
        tbl = kids[41]
        for tr in tbl.findall(qn("w:tr"))[1:]:
            tcs = tr.findall(qn("w:tc"))
            if len(tcs) == 3 and text_of(tcs[0]).strip() not in ("Discussion Committee", "PGS-Reviewer", "Head of PGS"):
                for tc in tcs[:2]:
                    for tt in tc.iter(qn("w:t")):
                        tt.text = ""

    def replace_block(self, first, last, new):
        """Replace body children first..last (inclusive elements) by new list."""
        parent = first.getparent()
        idx = list(parent).index(first)
        cur = first
        doomed = []
        while True:
            doomed.append(cur)
            if cur is last:
                break
            cur = cur.getnext()
        for e in doomed:
            parent.remove(e)
        for k, e in enumerate(new):
            parent.insert(idx + k, e)

    def abstract_and_ack(self):
        kids = list(self.body)
        ab = [strip_prefix(clean_runs(p), 0) for p in pandoc_paragraphs(self.conv / "prep/abstract.docx")]
        ab = [r for r in ab if r]
        new = [self.para(r, after=200, first_line=720, line=360) for r in ab]
        new[-1].find(qn("w:pPr")).find(qn("w:ind")).set(qn("w:firstLine"), "0")
        self.replace_block(kids[47], kids[47], new)
        kids = list(self.body)
        ack_i = next(i for i, k in enumerate(kids) if text_of(k).strip().startswith("Special acknowledgment"))
        ac = [clean_runs(p) for p in pandoc_paragraphs(self.conv / "prep/acknowledgments.docx")]
        ac = [r for r in ac if r]
        new = [self.para(r, after=200, first_line=720, line=360) for r in ac]
        self.replace_block(kids[ack_i], kids[ack_i + 1], new)

    def toc_fields(self):
        """Empty the cached TOC / figure / table lists; the fields are rebuilt on update."""
        body = self.body
        sdt = body.find(qn("w:sdt"))
        content = sdt.find(qn("w:sdtContent"))
        ps = content.findall(qn("w:p"))
        instr = None
        for p in ps:
            for it in p.iter(qn("w:instrText")):
                if it.text and "TOC" in it.text:
                    instr = it.text
                    break
            if instr:
                break
        keep = ps[0]
        for p in ps[1:]:
            content.remove(p)
        content.append(self.field_paragraphs(instr, "TOC1")[0])
        content.append(self.field_paragraphs(instr, "TOC1")[1])
        for style_key in ("my figures ibrahim", "my tables ibrahim"):
            kids = list(body)
            start = next(i for i, k in enumerate(kids)
                         if any(style_key in (it.text or "") for it in k.iter(qn("w:instrText"))))
            depth, end = 0, None
            for j in range(start, len(kids)):
                for fc in kids[j].iter(qn("w:fldChar")):
                    typ = fc.get(qn("w:fldCharType"))
                    depth += 1 if typ == "begin" else (-1 if typ == "end" else 0)
                    if depth == 0 and typ == "end":
                        end = j
                        break
                if end is not None:
                    break
            instr = next(it.text for it in kids[start].iter(qn("w:instrText")) if style_key in (it.text or ""))
            self.replace_block(kids[start], kids[end], self.field_paragraphs(instr, "TableofFigures"))

    def fill_indexes(self):
        data = json.loads(Path(self.indexes).read_text())
        toc = next(x["text"] for x in data if x["name"].startswith("Table of Contents"))
        figs = [x["text"] for x in data if x["name"].startswith("Table of Figures")]

        def entries(text):
            out = []
            for line in text.split("\n"):
                line = line.strip("\t ")
                if not line or line == "Contents" or "\t" not in line:
                    continue
                title, page = line.rsplit("\t", 1)
                out.append((title.replace("\t", " ").strip(), page.strip()))
            return out

        def level(t):
            if re.match(r"^([0-9]+|[A-Z])\.[0-9]+\.[0-9]+ ", t):
                return "TOC3"
            if re.match(r"^([0-9]+|[A-Z])\.[0-9]+ ", t):
                return "TOC2"
            return "TOC1"
        fields = []
        for p in self.d.element.body.iter(qn("w:p")):
            for it in p.iter(qn("w:instrText")):
                if it.text and it.text.strip().startswith("TOC"):
                    fields.append((p, it.text))
        for p, instr in fields:
            is_toc = "\\o" in instr
            items = entries(toc) if is_toc else entries(figs.pop(0) if "figures" in instr else figs.pop(-1))
            if not is_toc and "figures" in instr:
                pass
            sep = [r for r in p.findall(qn("w:r")) if r.find(qn("w:fldChar")) is not None
                   and r.find(qn("w:fldChar")).get(qn("w:fldCharType")) == "separate"][0]
            for r in list(p.findall(qn("w:r")))[list(p.findall(qn("w:r"))).index(sep) + 1:]:
                p.remove(r)
            style = p.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val"))
            prev = p
            for k, (title, page) in enumerate(items):
                st = level(title) if is_toc else style
                runs = [run(title), tab_run(), run(page)]
                if k == 0:
                    tabs = sub(p.find(qn("w:pPr")), "w:tabs")
                    sub(tabs, "w:tab", **{"w:val": "right", "w:leader": "dot", "w:pos": 9961})
                    p.find(qn("w:pPr")).find(qn("w:pStyle")).set(qn("w:val"), st)
                    for r in runs:
                        p.append(r)
                    continue
                q = self.para(runs, style=st, after=60 if not is_toc else 100, tabs=[("right", 9961)])
                q.find(qn("w:pPr")).find(qn("w:tabs")).find(qn("w:tab")).set(qn("w:leader"), "dot")
                prev.addnext(q)
                prev = q

    def field_paragraphs(self, instr, style):
        p1 = el("w:p")
        ppr = sub(p1, "w:pPr")
        sub(ppr, "w:pStyle", **{"w:val": style})
        r = sub(p1, "w:r")
        sub(r, "w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"})
        r = sub(p1, "w:r")
        it = sub(r, "w:instrText")
        it.text = instr
        it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r = sub(p1, "w:r")
        sub(r, "w:fldChar", **{"w:fldCharType": "separate"})
        p1.append(run(" "))
        p2 = el("w:p")
        ppr = sub(p2, "w:pPr")
        sub(ppr, "w:pStyle", **{"w:val": style})
        r = sub(p2, "w:r")
        sub(r, "w:fldChar", **{"w:fldCharType": "end"})
        return [p1, p2]

    def abbreviations(self):
        kids = list(self.body)
        start = next(i for i, k in enumerate(kids) if text_of(k).strip() == "Abbreviations")
        end = next(i for i, k in enumerate(kids) if text_of(k).strip().startswith("WAVE"))
        base = copy.deepcopy(self.proto["abbr"].find(qn("w:r")).find(qn("w:rPr")))
        new = []
        for block in self.side["abbreviations"]:
            for a, b in block["rows"]:
                runs = set_size(self.snippet(a), 24) + [tab_run()] + set_size(self.snippet(b), 24)
                new.append(self.para(runs, style="Normal", after=60, tabs=[("left", 1900)], left=1900,
                                     hanging=1900, line=276))
        new.append(self.page_break())
        h = copy.deepcopy(kids[next(i for i, k in enumerate(kids) if text_of(k).strip() == "List of Abbreviations")])
        for b in h.findall(qn("w:bookmarkStart")) + h.findall(qn("w:bookmarkEnd")):
            h.remove(b)
        self.replace_paragraph_text(h, "List of Symbols")
        new.append(h)
        for block in self.side["symbols"]:
            if block["heading"]:
                new.append(self.para(bold_runs(set_size(self.snippet(block["heading"]), 24)), style="Normal",
                                     after=120, before=200, first_line=0, keep_next=True))
            for a, b in block["rows"]:
                runs = set_size(self.snippet(a), 24) + [tab_run()] + set_size(self.snippet(b), 24)
                new.append(self.para(runs, style="Normal", after=60, tabs=[("left", 1900)], left=1900,
                                     hanging=1900, line=276))
        self.replace_block(kids[start + 1], kids[end], new)

    def back_matter(self):
        out = []
        h = copy.deepcopy(self.proto["h1"])
        for b in h.findall(qn("w:bookmarkStart")) + h.findall(qn("w:bookmarkEnd")):
            h.remove(b)
        out.append(self.simple_h1("References"))
        refs = pandoc_paragraphs(self.conv / "refs.docx")
        bib = [p for p in refs if p.find(qn("w:pPr")) is not None and
               p.find(qn("w:pPr")).find(qn("w:pStyle")) is not None and
               p.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == "Bibliography"]
        order = (self.conv / "refs_order.txt").read_text().split()
        assert len(bib) == len(order)
        for k, p in enumerate(bib):
            runs = clean_runs(p)
            assert text_of(p).startswith(f"[{k + 1}]"), text_of(p)[:20]
            if len(runs) > 2 and text_of(runs[1]) == " " and runs[2].find(qn("w:tab")) is not None:
                del runs[1]
            out.append(self.para(runs, style="ListParagraph", left=567, hanging=567, jc="both", after=120,
                                 tabs=[("left", 567)], line=276))
        return out

    def simple_h1(self, text, label=None):
        p = el("w:p")
        ppr = sub(p, "w:pPr")
        sub(ppr, "w:pStyle", **{"w:val": "Heading1"})
        sub(ppr, "w:pageBreakBefore")
        if label:
            p.append(run(label))
            r = el("w:r")
            sub(r, "w:br")
            p.append(r)
        p.append(run(text))
        return p

    def publications(self):
        out = [self.simple_h1("Publications")]
        for k, sid in enumerate(self.side["publications"]):
            runs = self.snippet(sid)
            out.append(self.para([run(f"[P{k + 1}]"), tab_run()] + runs, style="ListParagraph", left=709, hanging=709,
                                 jc="both", after=160, tabs=[("left", 709)], line=276))
        return out

    def arabic_summary(self):
        src = (THESIS / "arabic/arabic_summary.tex").read_text()
        body = src.split("\\thispagestyle{empty}", 1)[1].split("\\clearpage", 1)[0]
        body = re.sub(r"\\foreignlanguage\{english\}\{([^}]*)\}", r"\1", body)
        body = re.sub(r"\\vspace\{[^}]*\}|\\noindent", "", body)
        paras = [x.strip() for x in re.split(r"\n\s*\n", body) if x.strip()]
        kids = list(self.body)
        i = next(i for i, k in enumerate(kids) if text_of(k).strip() == "ملخص الرسالة")
        proto = kids[i + 1]
        new = []
        for t in paras:
            p = copy.deepcopy(proto)
            for r in p.findall(qn("w:r")):
                p.remove(r)
            base = proto.find(qn("w:r")).find(qn("w:rPr"))
            m = re.match(r"\\textbf\{([^}]*)\}(.*)", t, re.S)
            if m:
                p.append(run(m.group(1), bold=True, base_rpr=base))
                p.append(run(m.group(2).strip() and " " + m.group(2).strip(), base_rpr=base))
            else:
                p.append(run(re.sub(r"\s+", " ", t), base_rpr=base))
            new.append(p)
        self.replace_block(proto, proto, new)

    def schema_order(self):
        for rpr in self.d.element.iter(qn("w:rPr")):
            kids = list(rpr)
            kids.sort(key=lambda e: RPR_ORDER.get(etree.QName(e).localname, 99))
            for e in kids:
                rpr.append(e)
        for ppr in self.d.element.iter(qn("w:pPr")):
            kids = list(ppr)
            kids.sort(key=lambda e: PPR_ORDER.get(etree.QName(e).localname, 99))
            for e in kids:
                ppr.append(e)

    def drop_orphan_parts(self):
        part = self.d.part
        xml = etree.tostring(self.d.element, encoding="unicode")
        used = set(re.findall(r'r:(?:embed|id|link|pict)="([^"]+)"', xml))
        kinds = ("/image", "/chart", "/oleObject", "/package", "/hyperlink", "/diagramData", "/diagramLayout",
                 "/diagramQuickStyle", "/diagramColors", "/diagramDrawing")
        for rid, rel in list(part.rels.items()):
            if rel.reltype.endswith(kinds) and rid not in used:
                part.drop_rel(rid)

    def clean_package(self):
        import datetime
        w14 = "http://schemas.microsoft.com/office/word/2010/wordml"
        for p in self.d.element.iter(qn("w:p")):
            for a in ("paraId", "textId"):
                p.attrib.pop(f"{{{w14}}}{a}", None)
        cp = self.d.core_properties
        now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
        cp.created = now
        cp.modified = now
        cp.revision = 1
        cp.last_printed = now
        for part in self.d.part.package.iter_parts():
            if part.partname == "/docProps/custom.xml":
                part._blob = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Properties xmlns="http://'
                              b'schemas.openxmlformats.org/officeDocument/2006/custom-properties" xmlns:vt="http://'
                              b'schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>')
            elif part.partname == "/docProps/app.xml":
                part._blob = re.sub(rb"<TotalTime>\d+</TotalTime>", b"<TotalTime>0</TotalTime>", part.blob)

    def rename_caption_styles(self):
        ren = {"myfiguresibrahim": ("thesisfigures", "thesis figures"),
               "mytablesibrahim": ("thesistables", "thesis tables"),
               "myfiguresibrahimChar": ("thesisfiguresChar", "thesis figures Char"),
               "mytablesibrahimChar": ("thesistablesChar", "thesis tables Char")}
        st = self.d.styles.element
        for s in st.findall(qn("w:style")):
            sid = s.get(qn("w:styleId"))
            if sid in ren:
                s.set(qn("w:styleId"), ren[sid][0])
                s.find(qn("w:name")).set(qn("w:val"), ren[sid][1])
            for tag in ("w:basedOn", "w:link", "w:next"):
                e = s.find(qn(tag))
                if e is not None and e.get(qn("w:val")) in ren:
                    e.set(qn("w:val"), ren[e.get(qn("w:val"))][0])
        for e in self.d.element.iter(qn("w:pStyle"), qn("w:rStyle")):
            if e.get(qn("w:val")) in ren:
                e.set(qn("w:val"), ren[e.get(qn("w:val"))][0])
        for it in self.d.element.iter(qn("w:instrText")):
            if it.text:
                it.text = it.text.replace("my figures ibrahim", "thesis figures").replace("my tables ibrahim", "thesis tables")

    def front_layout(self):
        kids = list(self.body)
        start = next(i for i, k in enumerate(kids) if text_of(k).strip() == "Abstract")
        end = next(i for i, k in enumerate(kids) if i > start and k.tag == qn("w:p")
                   and k.find(qn("w:pPr")) is not None and k.find(qn("w:pPr")).find(qn("w:sectPr")) is not None)
        heads = {"Acknowledgments", "List of Figures", "List of Tables", "List of Abbreviations", "List of Symbols"}
        for k in kids[start:end]:
            if k.tag != qn("w:p"):
                continue
            t = text_of(k).strip()
            ppr = k.find(qn("w:pPr"))
            if t in heads and ppr is not None:
                if ppr.find(qn("w:pageBreakBefore")) is None:
                    ppr.append(el("w:pageBreakBefore"))
                continue
            busy = (k.findall(".//" + qn("w:drawing")) or k.findall(".//" + qn("w:fldChar"))
                    or k.findall(".//" + qn("w:instrText")) or k.findall(".//" + f"{{{M_NS}}}oMath"))
            if not t and not busy:
                self.body.remove(k)
        sdt = self.body.find(qn("w:sdt"))
        title = sdt.find(qn("w:sdtContent")).find(qn("w:p"))
        title.find(qn("w:pPr")).append(el("w:pageBreakBefore"))
        kids = list(self.body)
        sig = next(k for k in kids if text_of(k).strip() == AUTHOR and k.find(qn("w:pPr")) is not None
                   and k.getprevious() is not None and "family" in text_of(k.getprevious()))
        sp = sig.find(qn("w:pPr")).find(qn("w:spacing"))
        if sp is None:
            sp = sub(sig.find(qn("w:pPr")), "w:spacing")
        sp.set(qn("w:before"), "720")
        # Arabic title page: the two-line title takes the place of two spacer lines
        kids = list(self.body)
        date = next(i for i, k in enumerate(kids) if text_of(k).strip() == f"القاهرة {YEAR}")
        removed = 0
        j = date - 1
        while removed < 2 and not text_of(kids[j]).strip():
            self.body.remove(kids[j])
            removed += 1
            j -= 1

    def settings(self):
        s = self.d.settings.element
        if s.find(qn("w:updateFields")) is None:
            u = el("w:updateFields", **{"w:val": "true"})
            s.insert(0, u)

    def build(self, out):
        self.front()
        self.abstract_and_ack()
        self.toc_fields()
        self.abbreviations()
        kids = list(self.body)
        first = next(k for k in kids if k.tag == qn("w:p") and text_of(k).replace("\n", " ").strip().startswith("Chapter 1")
                     and k.find(qn("w:pPr")) is not None and k.find(qn("w:pPr")).find(qn("w:pStyle")) is not None
                     and k.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == "Heading1")
        sect_i = next(i for i, k in enumerate(kids) if text_of(k).strip() == "ملخص الرسالة") - 1
        last = kids[sect_i - 1]
        content = self.content()
        app_i = next(i for i, e in enumerate(content) if text_of(e).startswith("Appendix A"))
        chapters, appendix = content[:app_i], content[app_i:]
        new = chapters + self.back_matter() + appendix + self.publications()
        self.replace_block(first, last, new)
        self.arabic_summary()
        self.front_layout()
        if self.indexes:
            self.fill_indexes()
        self.rename_caption_styles()
        self.schema_order()
        self.drop_orphan_parts()
        self.clean_package()
        self.settings()
        self.d.core_properties.author = AUTHOR
        self.d.core_properties.title = TITLE
        self.d.core_properties.last_modified_by = AUTHOR
        self.d.core_properties.comments = ""
        self.d.save(out)


if __name__ == "__main__":
    Builder(sys.argv[1], sys.argv[2], sys.argv[4] if len(sys.argv) > 4 else None).build(sys.argv[3])
